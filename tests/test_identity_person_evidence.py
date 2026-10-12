"""Row references, private address evidence and keyed fingerprints, on a fictional world (tests/identity_person_fixtures).

The raw workbooks here hold fictional contributor streets on purpose: every test that writes anything also proves the
streets never left memory.
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

import identity_person as cli  # noqa: E402
import identity_person_evidence as ev  # noqa: E402
import identity_person_key as ik  # noqa: E402
from tests.identity_person_fixtures import (  # noqa: E402
    COUNCIL, SOURCE, assert_no_street, base_rows, build_world, every_text,
)

KEY = ik.Key(1, bytes(range(32)))
KEY_2 = ik.Key(2, bytes(range(1, 33)))


@pytest.fixture
def world(tmp_path, monkeypatch):
    w = build_world(tmp_path)
    monkeypatch.setenv(ik.ENV, str(w.private))
    return w


def _collect(w, key=KEY):
    return ev.collect(ev.Inputs(w.ledger_dir, w.raw_root, w.export_dir, w.decisions), key)



def _one(evidence, tran_id, filer="1400001"):
    [row] = [r for r in evidence.rows if r.tran_id == tran_id and r.filer_id == filer][:1]
    return row


# --- address canonicalization (in memory only) -------------------------------------------------------------------

@pytest.mark.parametrize("a, b", [
    (("12 Sample Street", "Apt. 3", "94901-1234"), ("12 SAMPLE ST #3", None, 94901)),
    (("12 Sample Street", "", "94901"), ("12  sample   st.", "N/A", "949011234")),
    (("400 North Example Avenue", None, "94903"), ("400 N. Example Ave", "", "94903")),
    (("P.O. Box 12", None, "94901"), ("PO BOX 12", None, "94901")),
    (("Post Office Box 12", None, "94901"), ("POBOX 12", None, "94901")),
    (("５ Placeholder Lane", None, "94901"), ("5 Placeholder Ln", None, "94901")),  # full-width digit, NFKC
    (("9 Imaginary Court Suite 4", None, "94901"), ("9 Imaginary Ct Ste 4", None, "94901")),
])
def test_equivalent_addresses_canonicalize_alike(a, b):
    one, two = ev.canonical_address(*a), ev.canonical_address(*b)
    assert (one.street, one.zip5, one.reason) == (two.street, two.zip5, two.reason)
    assert one.street is not None and ev.address_hash(KEY, one) == ev.address_hash(KEY, two)


@pytest.mark.parametrize("adr1, reason", [
    ("N/A", "placeholder_street"), ("n.a.", "placeholder_street"), ("NONE", "placeholder_street"),
    ("Information Requested", "placeholder_street"), ("Best Efforts", "placeholder_street"),
    ("--", "placeholder_street"), ("xxx", "placeholder_street"), ("000", "placeholder_street"),
    ("", "missing_street"), (None, "missing_street"), (12, "non_text_street"),
])
def test_placeholder_and_missing_streets_get_no_hash(adr1, reason):
    address = ev.canonical_address(adr1, None, "94901")
    assert (address.street, address.reason) == (None, reason)


def test_placeholders_are_detected_before_punctuation_is_stripped():
    # "N/A" would become "N A" after stripping; detection on the raw text catches it, and "N A" itself is not a
    # placeholder pattern that would also swallow a real (if odd) street like "N A Example Rd".
    assert ev.canonical_address("N/A", None, "94901").reason == "placeholder_street"
    assert ev.canonical_address("5 N A Example Rd", None, "94901").street is not None


@pytest.mark.parametrize("zip_cell", ["9490", "ABCDE", "00000", "94901-12", None, "", 9490])
def test_an_invalid_or_missing_zip_gives_no_hash(zip_cell):
    address = ev.canonical_address("12 Sample Street", None, zip_cell)
    assert address.street is None and address.reason in ("invalid_zip", "missing_zip")


def test_a_placeholder_second_line_is_dropped():
    assert ev.canonical_address("12 Sample Street", "N/A", "94901").street == ev.canonical_address(
        "12 Sample Street", None, "94901").street


def test_a_cosmetic_change_keeps_the_cluster_but_changes_the_evidence(tmp_path):
    before = _one(_collect(build_world(tmp_path / "one")), "A-1")
    rows = base_rows()
    rows["A-Contributions"][0]["Tran_Adr1"] = "12 Sample St"
    after = _one(_collect(build_world(tmp_path / "two", rows=rows)), "A-1")
    assert before.address_hash == after.address_hash
    assert ev.diff_fields(before.fp, after.fp) == ["address"]


def test_address_hashes_are_keyed_versioned_and_domain_separated():
    address = ev.canonical_address("12 Sample Street", None, "94901")
    h1, h2 = ev.address_hash(KEY, address), ev.address_hash(KEY_2, address)
    assert len(h1) == 64 and h1 != h2
    assert h1 != ev.address_hash(ik.Key(1, bytes(range(1, 33))), address)
    assert ev.address_hash(KEY, ev.canonical_address("N/A", None, "94901")) is None


# --- row references ----------------------------------------------------------------------------------------------

VECTOR_KEY = ["example-cf", "1400001", "2024-01-01", "2024-01-20", "2024-02-01", "000", "A", "T-1"]


def test_row_id_is_pinned_and_holds_no_file_sheet_row_or_sha():
    assert ev.row_id(VECTOR_KEY) == "r-5d2c5d5869b3cf55b725"


def test_address_hash_and_fingerprint_vectors_are_pinned():
    address = ev.canonical_address("12 Sample Street", "Apt. 3", "94901-1234")
    assert ev.address_hash(KEY, address) == "e508850463192c6779f355d8af85d442c258781c219be9fbdea3d260c725714c"
    values = {"name": ["Doe", "Pat", None, None], "entity": "IND", "filing": "filing-0", "date": "2024-01-05",
              "amount": "5.00", "memo": [None, None], "flow": ["moneyflow-1", True, "retained", "person-doe-pat"],
              "occupation": "Teacher", "employer": None, "city": "Exampleton", "state": "CA", "zip": "94901",
              "address": ["12 Sample Street", None, "94901"]}
    assert ev.fingerprint(KEY, ev.row_id(VECTOR_KEY), values) == (
        "fp1:k1:9c60bf211bfb5213606c36fe8a13f632a4963691de6594610b04ff47b4abc2a2:0f99e337.fa74f6b6.06dcfbba."
        "33c53a40.f7872d62.fe277c10.91576711.6de9012a.2e16eb87.30dcdfac.31527889.efcec813.dd647736")


def test_row_ids_and_fingerprints_survive_a_redownload_that_reorders_rows(tmp_path):
    first = _collect(build_world(tmp_path / "one"))
    rows = base_rows()
    rows["A-Contributions"] = list(reversed(rows["A-Contributions"]))
    second = _collect(build_world(tmp_path / "two", rows=rows))
    assert {r.row_id: r.fp for r in first.rows} == {r.row_id: r.fp for r in second.rows}
    assert {r.row_id: r.address_hash for r in first.rows} == {r.row_id: r.address_hash for r in second.rows}


def test_scope_is_counted_individual_contributions_from_person_buckets(world):
    evidence = _collect(world)
    assert len(evidence.rows) == 14  # every A row; the Vale flow has two physical rows
    assert {r.bucket for r in evidence.rows} >= {"person-quill-avery", "person-vale-morgan"}
    vale = [r for r in evidence.rows if r.bucket == "person-vale-morgan"]
    assert len(vale) == 2 and len({r.row_id for r in vale}) == 2 and len({r.moneyflow_id for r in vale}) == 1


def test_a_row_key_shared_by_two_rows_is_ambiguous_and_out_of_scope(tmp_path):
    rows = base_rows()
    twin = dict(rows["A-Contributions"][0], Tran_Amt1=5, Tran_Amt2=5, Tran_NamL="Twin")
    rows["A-Contributions"].append(twin)
    evidence = _collect(build_world(tmp_path, rows=rows))
    assert evidence.counts["ambiguous_row_key"] >= 1
    assert not [r for r in evidence.rows if r.tran_id == "A-1"]


def test_treasurer_and_intermediary_streets_are_never_read(tmp_path):
    first = _collect(build_world(tmp_path / "one"))
    rows = base_rows()
    for row in rows["A-Contributions"]:
        row.update(Tres_Adr1="3 Other Treasurer Way", Intr_Adr1="4 Other Intermediary Way")
    second = _collect(build_world(tmp_path / "two", rows=rows))
    assert [(r.row_id, r.fp, r.address_hash) for r in first.rows] == [
        (r.row_id, r.fp, r.address_hash) for r in second.rows]


# --- fingerprints -------------------------------------------------------------------------------------------------

def test_an_address_only_change_changes_only_the_address_field(tmp_path):
    before = _one(_collect(build_world(tmp_path / "one")), "A-1")
    rows = base_rows()
    rows["A-Contributions"][0]["Tran_Adr1"] = "400 Example Avenue"
    after = _one(_collect(build_world(tmp_path / "two", rows=rows)), "A-1")
    assert before.row_id == after.row_id and before.fp != after.fp
    assert ev.diff_fields(before.fp, after.fp) == ["address"]


def test_an_address_change_without_a_hash_still_changes_the_fingerprint(tmp_path):
    rows = base_rows()
    rows["A-Contributions"][0]["Tran_Zip4"] = "9490"
    before = _one(_collect(build_world(tmp_path / "one", rows=rows)), "A-1")
    rows["A-Contributions"][0]["Tran_Adr1"] = "400 Example Avenue"
    after = _one(_collect(build_world(tmp_path / "two", rows=rows)), "A-1")
    assert before.address_hash is None and after.address_hash is None
    assert ev.diff_fields(before.fp, after.fp) == ["address"]


def test_an_employer_change_is_named_in_the_diff(tmp_path):
    before = _one(_collect(build_world(tmp_path / "one")), "A-1")
    rows = base_rows()
    rows["A-Contributions"][0]["Tran_Emp"] = "Example Water District"
    after = _one(_collect(build_world(tmp_path / "two", rows=rows)), "A-1")
    assert ev.diff_fields(before.fp, after.fp) == ["employer"]


def test_fingerprints_are_keyed_and_versioned(world):
    a, b = _collect(world, KEY), _collect(world, KEY_2)
    one, two = a.rows[0], next(r for r in b.rows if r.row_id == a.rows[0].row_id)
    assert one.fp.startswith("fp1:k1:") and two.fp.startswith("fp1:k2:")
    assert ev.diff_fields(one.fp, two.fp) == ev.FP_FIELDS  # re-keyed: every field differs


def test_fingerprint_shape_is_a_full_row_hmac_plus_short_field_digests(world):
    fp = _collect(world).rows[0].fp
    prefix, version, full, fields = fp.split(":")
    assert (prefix, version, len(full)) == ("fp1", "k1", 64)
    assert [len(d) for d in fields.split(".")] == [8] * len(ev.FP_FIELDS)


def test_equal_values_in_different_rows_do_not_share_field_digests(world):
    rows = _collect(world).rows
    a, b = (r for r in rows if r.tran_id in ("A-1", "A-2"))
    assert a.fp.split(":")[3].split(".")[ev.FP_FIELDS.index("city")] != \
        b.fp.split(":")[3].split(".")[ev.FP_FIELDS.index("city")]


# --- inputs -------------------------------------------------------------------------------------------------------

def test_a_ledger_row_that_does_not_match_its_workbook_row_fails_by_reference(world):
    path = world.ledger_dir / SOURCE / "ledger.jsonl"
    rows = [json.loads(line) for line in path.read_text().splitlines()]
    a = [r for r in rows if r.get("schedule") == "A"]
    a[0]["row_ref"]["row"], a[1]["row_ref"]["row"] = a[1]["row_ref"]["row"], a[0]["row_ref"]["row"]
    path.write_text("".join(json.dumps(r, sort_keys=True) + "\n" for r in rows))
    with pytest.raises(ev.InputError, match=r"r-[0-9a-f]{20}: workbook row does not match the ledger") as info:
        _collect(world)
    assert_no_street(str(info.value))


def test_the_raw_workbook_must_be_the_newest_capture_of_its_year(world):
    newer = world.raw_root / SOURCE / "2026-11-01" / "2024.zip"
    newer.parent.mkdir(parents=True)
    newer.write_bytes(b"PK not the same bytes")
    with pytest.raises(ev.InputError, match="not the newest raw capture"):
        _collect(world)


def test_a_ledger_input_missing_from_the_raw_root_fails(world):
    for p in (world.raw_root / SOURCE).rglob("*.zip"):
        p.unlink()
    with pytest.raises(ev.InputError, match="no raw workbook"):
        _collect(world)


def test_a_corrupt_workbook_fails_without_echoing_its_contents(world):
    [path] = list((world.raw_root / SOURCE).rglob("*.zip"))
    data = path.read_bytes()
    path.write_bytes(data[:200] + b"12 Sample Street" + data[200:])  # sha changes too: make the ledger agree
    manifest = world.ledger_dir / SOURCE / "manifest.json"
    pinned = json.loads(manifest.read_text())
    pinned["inputs"][0]["sha256"] = __import__("hashlib").sha256(path.read_bytes()).hexdigest()
    manifest.write_text(json.dumps(pinned))
    with pytest.raises(ev.InputError, match="not a readable workbook") as info:
        _collect(world)
    assert_no_street(str(info.value))


def _date_styled(world, column: str, value) -> None:
    """Rewrite the raw workbook so the first A row's `column` holds `value` styled as a date (openpyxl warns when
    such a number cannot become a date, quoting the value)."""
    import hashlib
    import zipfile

    import openpyxl

    [path] = list((world.raw_root / SOURCE).rglob("*.zip"))
    with zipfile.ZipFile(path) as zf:
        [inner] = zf.namelist()
        wb = openpyxl.load_workbook(zf.open(inner))
    ws = wb["A-Contributions"]
    header = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=header.index(column) + 1, value=value)
    cell.number_format = "yyyy-mm-dd"
    xlsx = path.with_suffix(".xlsx")
    wb.save(xlsx)
    with zipfile.ZipFile(path, "w") as zf:
        zf.write(xlsx, inner)
    xlsx.unlink()
    manifest = world.ledger_dir / SOURCE / "manifest.json"
    pinned = json.loads(manifest.read_text())
    pinned["inputs"][0]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest.write_text(json.dumps(pinned))


def test_parser_warnings_that_quote_a_street_cell_never_surface(world, capsys, recwarn):
    _date_styled(world, "Tran_Adr1", 98765432123)
    cli.main(["init-key"])
    code, said = _build(world, capsys)
    assert code == 0, said
    assert "98765432123" not in said and not [w for w in recwarn if "98765432123" in str(w.message)]
    assert "98765432123" not in every_text([world.root], skip=world.inputs)


def test_a_workbook_the_parser_cannot_read_losslessly_puts_all_its_rows_out_of_scope(world):
    # openpyxl parses ahead of the rows it yields, so a warning cannot be pinned to one row: the whole workbook
    # is lossy, and its rows get no evidence at all (two different streets could otherwise hash alike).
    _date_styled(world, "Tran_Adr1", 98765432123)
    evidence = _collect(world)
    assert evidence.counts["lossy_workbook_rows"] == 14
    assert evidence.rows == []


def test_a_non_text_street_cell_is_a_reason_not_a_crash(tmp_path):
    rows = base_rows()
    rows["A-Contributions"][0]["Tran_Adr1"] = 12345
    row = _one(_collect(build_world(tmp_path, rows=rows)), "A-1")
    assert (row.address_hash, row.address_reason) == (None, "non_text_street")


# --- the build command --------------------------------------------------------------------------------------------

def _build(w, capsys, *extra, as_of="2026-10-12", out="test"):
    code = cli.main([*w.build_args(as_of=as_of, out=out), *extra])
    captured = capsys.readouterr()
    return code, captured.out + captured.err


def test_build_without_a_key_fails_closed_and_writes_nothing(world, capsys):
    code, said = _build(world, capsys)
    assert code == 1 and "run init-key" in said
    assert not world.private.exists()


def test_build_writes_the_manifest_and_address_evidence_only_in_the_private_dir(world, capsys):
    assert cli.main(["init-key"]) == 0
    code, said = _build(world, capsys)
    assert code == 0, said
    out = world.private / "builds" / "test"
    manifest = json.loads((out / "manifest.json").read_text())
    assert manifest["as_of"] == "2026-10-12" and manifest["key_version"] == 1
    roles = {i["role"] for i in manifest["inputs"]}
    assert roles >= {"ledger", "raw_workbook", "export_nodes", "export_edges", "decisions"}
    assert all(len(i["sha256"]) == 64 for i in manifest["inputs"])
    evidence = [json.loads(line) for line in (out / "address-evidence.jsonl").read_text().splitlines()]
    hashes = {e["address_hash"] for e in evidence if e["address_hash"]}
    assert hashes and all(len(h) == 64 for h in hashes)
    outside = every_text([world.root], skip=world.inputs, exclude=[world.private]) + said
    assert not any(h in outside for h in hashes)
    assert_no_street(every_text([world.root], skip=world.inputs) + said)
    assert oct(os.stat(out / "address-evidence.jsonl").st_mode & 0o777) == "0o600"


def test_two_builds_over_the_same_inputs_are_byte_identical(world, capsys):
    cli.main(["init-key"])
    _build(world, capsys, out="one")
    _build(world, capsys, out="two")
    one, two = world.private / "builds" / "one", world.private / "builds" / "two"
    assert sorted(p.name for p in one.iterdir()) == sorted(p.name for p in two.iterdir())
    for p in one.iterdir():
        assert p.read_bytes() == (two / p.name).read_bytes(), p.name


def test_a_rebuild_replaces_the_previous_output(world, capsys):
    cli.main(["init-key"])
    _build(world, capsys)
    stray = world.private / "builds" / "test" / "stray.txt"
    stray.write_text("old")
    assert _build(world, capsys)[0] == 0
    assert not stray.exists()


def test_a_build_named_like_the_key_lands_under_builds_and_leaves_the_key_alone(world, capsys):
    cli.main(["init-key"])
    secret = ik.load_key(world.private).secret
    assert _build(world, capsys, out="key-v1")[0] == 0
    assert (world.private / "builds" / "key-v1" / "manifest.json").is_file()
    assert ik.load_key(world.private).secret == secret


@pytest.mark.parametrize("out", ["../escape", "/tmp/elsewhere", "a/b", "..", ".", ""])
def test_build_output_is_a_plain_name_under_builds(world, capsys, out):
    cli.main(["init-key"])
    secret = ik.load_key(world.private).secret
    code, said = _build(world, capsys, out=out)
    assert code == 1 and "plain name" in said
    assert ik.load_key(world.private).secret == secret


def test_build_refuses_an_output_inside_a_nested_checkout(world, capsys):
    cli.main(["init-key"])
    (world.private / "builds" / ".git").mkdir(parents=True)
    code, said = _build(world, capsys)
    assert code == 1 and "git checkout" in said
    assert not (world.private / "builds" / "test").exists()


def test_build_replaces_only_a_previous_build(world, capsys):
    cli.main(["init-key"])
    (world.private / "builds" / "test").mkdir(parents=True)
    (world.private / "builds" / "test" / "notes.txt").write_text("mine")
    code, said = _build(world, capsys)
    assert code == 1 and "not a previous build" in said
    assert (world.private / "builds" / "test" / "notes.txt").read_text() == "mine"


def test_unexpected_errors_print_only_their_type(world, capsys, monkeypatch):
    cli.main(["init-key"])

    def boom(*_args, **_kwargs):
        raise RuntimeError("12 Sample Street")
    monkeypatch.setattr(ev, "collect", boom)
    code, said = _build(world, capsys)
    assert code == 2 and "internal error: RuntimeError (details suppressed)" in said
    assert_no_street(said)


def test_init_key_prints_only_the_path_and_version(world, capsys):
    assert cli.main(["init-key"]) == 0
    said = capsys.readouterr().out
    secret = ik.load_key(world.private).secret
    assert "key-v1" in said and secret.hex() not in said and secret.hex().upper() not in said


def test_init_key_refuses_a_version_already_used_by_the_decision_file(world, capsys):
    world.decisions.write_text(json.dumps({"fp": {"r-x": "fp1:k3:" + "0" * 64 + ":x"}}) + "\n")
    code = cli.main(["init-key", "--decisions", str(world.decisions)])
    assert code == 1 and "version 4 or higher" in capsys.readouterr().err
    assert cli.main(["init-key", "--decisions", str(world.decisions), "--version", "4"]) == 0


def test_council_fixture_is_what_the_tests_think_it_is():
    assert COUNCIL["Filer_ID"] == "1400001"


def test_the_cli_streams_never_carry_a_street_or_hash_on_success_or_failure(world):
    import subprocess

    env = {**os.environ, ik.ENV: str(world.private)}
    script = str(Path(__file__).resolve().parent.parent / "scripts" / "identity_person.py")

    def run(*args):
        done = subprocess.run([sys.executable, script, *args], capture_output=True, text=True, env=env)
        return done.returncode, done.stdout + done.stderr

    streams = []
    streams.append(run(*world.build_args()))  # no key yet
    assert streams[-1][0] == 1
    streams.append(run("init-key"))
    streams.append(run(*world.build_args()))
    assert streams[-1][0] == 0, streams[-1][1]
    path = world.ledger_dir / SOURCE / "ledger.jsonl"
    rows = [json.loads(line) for line in path.read_text().splitlines()]
    a = [r for r in rows if r.get("schedule") == "A"]
    a[0]["row_ref"]["row"], a[1]["row_ref"]["row"] = a[1]["row_ref"]["row"], a[0]["row_ref"]["row"]
    path.write_text("".join(json.dumps(r, sort_keys=True) + "\n" for r in rows))
    streams.append(run(*world.build_args(out="broken")))
    assert streams[-1][0] == 1
    text = "\n".join(s for _, s in streams)
    assert "Traceback" not in text
    assert_no_street(text + every_text([world.root], skip=world.inputs))
    hashes = [json.loads(line)["address_hash"] for line in
              (world.private / "builds" / "test" / "address-evidence.jsonl").read_text().splitlines()]
    assert not any(h and h in text for h in hashes)
