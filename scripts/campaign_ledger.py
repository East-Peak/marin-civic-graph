"""Campaign-finance ledger for NetFile yearly exports: hashed inputs, every physical row, reconciled filings.

The ledger is operator-private. It keeps every row of every sheet with a disposition, groups rows into
filings, and reconciles itemized Schedule A/E rows against the filer's own Summary totals. The public graph
is emitted from it by normalize_campaign_finance.py; nothing here writes outside the output root.
"""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

FILING_HEADERS = ("Filer_ID", "Filer_NamL", "Report_Num", "Rpt_Date", "From_Date", "Thru_Date", "Form_Type")
_MEMO_HEADERS = ("Memo_Code", "Memo_RefNo")
# Every sheet of the NetFile export, with the columns the ledger reads from it. A workbook with a sheet not
# listed here fails: every sheet must be inventoried.
SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "A-Contributions": (*FILING_HEADERS, "Tran_ID", "Entity_Cd", "Tran_NamL", "Tran_NamF", "Tran_Date",
                        "Tran_Amt1", *_MEMO_HEADERS),
    "E-Expenditure": (*FILING_HEADERS, "Tran_ID", "Entity_Cd", "Payee_NamL", "Payee_NamF", "Expn_Date",
                      "Amount", *_MEMO_HEADERS),
    "Summary": (*FILING_HEADERS, "Line_Item", "Amount_A"),
    **{sheet: FILING_HEADERS for sheet in (
        "C-Contributions", "I-Contributions", "F496P3-Contributions", "F465P3-Expenditure",
        "F461P5-Expenditure", "D-Expenditure", "G-Expenditure", "F-Expenses", "B1-Loans", "B2-Loans",
        "H-Loans", "497", "496")},
}


class InputError(Exception):
    """An input is missing, unreadable, or not what the registry says it is."""


class UnsafeOutputError(Exception):
    """The output root would write into a protected location."""


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------

def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _check_workbook(path: Path) -> None:
    """Open the export and read every row of every sheet, so a corrupt sheet fails here, not later."""
    try:
        with zipfile.ZipFile(path) as zf:
            names = [n for n in zf.namelist() if n.endswith(".xlsx")]
            if len(names) != 1:
                raise InputError(f"{path.name}: expected one .xlsx inside, found {len(names)}")
            with zf.open(names[0]) as f:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                try:
                    _check_sheets(path.name, wb)
                finally:
                    wb.close()
    except InputError:
        raise
    except Exception as exc:  # zipfile, openpyxl and XML parse errors all mean "not a readable export"
        raise InputError(f"{path.name}: not a readable workbook export ({type(exc).__name__}: {exc})") from exc


def _check_sheets(name: str, wb) -> None:
    unknown = sorted(set(wb.sheetnames) - set(SHEET_HEADERS))
    missing = [s for s in SHEET_HEADERS if s not in wb.sheetnames]
    if unknown or missing:
        raise InputError(f"{name}: unknown sheet(s) {unknown}, missing sheet(s) {missing}")
    for sheet, required in SHEET_HEADERS.items():
        rows = wb[sheet].iter_rows(values_only=True)
        header = [h for h in next(rows, ()) if h is not None]
        absent = [h for h in required if h not in header]
        doubled = sorted({h for h in header if header.count(h) > 1})
        if absent or doubled:
            raise InputError(f"{name}!{sheet}: missing column(s) {absent}, duplicate column(s) {doubled}")
        for _ in rows:
            pass


def inventory_inputs(
    input_root: Path,
    source_id: str,
    capture_date: str,
    years: list[str],
    unavailable: list[dict],
) -> list[dict]:
    """Hash every yearly export of one capture; fail on anything unexpected.

    ``unavailable`` pins inputs known to be HTML pages rather than exports ({file, sha256}); they are recorded
    as unavailable coverage, never as empty years. Any other unreadable, missing or extra file fails.
    """
    capture = Path(input_root) / source_id / capture_date
    pins = {p["file"]: p["sha256"] for p in unavailable}
    expected = {f"{year}.zip": year for year in years}
    present = {p.name for p in capture.glob("*") if p.is_file()}
    extra = sorted(present - set(expected))
    if extra:
        raise InputError(f"{source_id}/{capture_date}: unexpected input(s) {', '.join(extra)}")
    entries = []
    for name, year in sorted(expected.items()):
        path = capture / name
        if not path.is_file():
            raise InputError(f"{source_id}/{capture_date}/{name}: missing")
        sha = _sha256(path)
        entry = {
            "path": f"{source_id}/{capture_date}/{name}",
            "source_id": source_id,
            "year": year,
            "bytes": path.stat().st_size,
            "sha256": sha,
        }
        if name in pins:
            if pins[name] != sha:
                raise InputError(f"{entry['path']}: pinned as unavailable with a different sha256")
            entry.update(coverage="unavailable", reason="html_page_not_export")
        else:
            _check_workbook(path)
            entry["coverage"] = "workbook"
        entries.append(entry)
    return entries


# ---------------------------------------------------------------------------
# Output root
# ---------------------------------------------------------------------------

def _protected_dirs(repo: Path) -> list[Path]:
    repo = Path(repo).resolve()
    dirs = [repo]
    for layer in ("normalized", "extracted"):  # symlinks into the private data repo
        link = repo / "data" / layer
        if link.exists():
            dirs.append(link.resolve().parent)
    for layer in ("normalized", "extracted", "exports", "ingest-runs"):
        dirs.append((repo / "data" / layer).resolve())
    return dirs


def _inside_git_checkout(path: Path) -> Path | None:
    for parent in (path, *path.parents):
        if (parent / ".git").exists():
            return parent
    return None


def resolve_output_root(output_root: Path, *repos: Path) -> Path:
    """Resolve symlinks and refuse any destination inside a git checkout or a protected data dir."""
    out = Path(output_root).expanduser().resolve()
    for repo in repos:
        for protected in _protected_dirs(repo):
            if out == protected or protected in out.parents:
                raise UnsafeOutputError(f"refusing to write into {protected}")
    checkout = _inside_git_checkout(out)
    if checkout is not None:
        raise UnsafeOutputError(f"refusing to write inside the git checkout {checkout}")
    if out.exists() and any(out.iterdir()):
        raise UnsafeOutputError(f"output root {out} is not empty")
    return out


# ---------------------------------------------------------------------------
# Ledger
# ---------------------------------------------------------------------------

class LedgerError(Exception):
    """An operator-supplied exception or version-evidence entry is incomplete or does not fit the data."""


@dataclass(frozen=True)
class Schedule:
    letter: str
    amount: str
    date: str
    name: str  # column prefix of the counterparty name: Tran_NamL / Payee_NamL
    reported: tuple[tuple[str, str], ...]  # (ledger key, column) of privately kept reported details


SCHEDULES = {
    "A-Contributions": Schedule("A", "Tran_Amt1", "Tran_Date", "Tran", (
        ("city", "Tran_City"), ("state", "Tran_State"), ("zip5", "Tran_Zip4"),
        ("employer", "Tran_Emp"), ("occupation", "Tran_Occ"))),
    "E-Expenditure": Schedule("E", "Amount", "Expn_Date", "Payee", (
        ("city", "Payee_City"), ("state", "Payee_State"), ("zip5", "Payee_Zip4"),
        ("expn_code", "Expn_Code"), ("expn_dscr", "Expn_Dscr"))),
}
ORACLE_LINE = "1"  # Summary Form_Type A/E Line_Item 1 Amount_A: the filer's itemized total for the period
CENT = Decimal("0.01")


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _iso(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return _text(value)


def _cents(value) -> Decimal:
    """Signed decimal cents; raises ValueError on anything that is not an exact cent amount."""
    if value is None or isinstance(value, bool) or (isinstance(value, str) and not value.strip()):
        raise ValueError("blank amount")
    try:
        amount = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f"not a number: {value!r}") from exc
    if not amount.is_finite() or amount != amount.quantize(CENT):
        raise ValueError(f"not a cent amount: {value!r}")
    return amount.quantize(CENT)


def _zip5(value) -> str | None:
    match = re.match(r"\s*(\d{5})", str(value or ""))
    return match.group(1) if match else None


def filer_key(filer_id: str, filer_name: str | None) -> str:
    """NetFile's Filer_ID, except 'Pending' (and any non-numeric id), which several new committees share."""
    return filer_id if filer_id.isdigit() else f"{filer_id}:{filer_name or ''}"


def _filing_key(source_id, filer_id, filer_name, report_num, rpt_date, from_date, thru_date) -> tuple:
    return (source_id, filer_key(filer_id, filer_name), _iso(from_date), _iso(thru_date), _iso(rpt_date),
            _text(report_num))


def _filing_id(key: tuple) -> str:
    return "filing-" + hashlib.sha256(json.dumps(key).encode()).hexdigest()[:16]


def _require(entry: dict, fields: tuple[str, ...], what: str) -> None:
    missing = [f for f in fields if not (isinstance(entry.get(f), str) and entry[f].strip())]
    if missing:
        raise LedgerError(f"{what} needs {', '.join(missing)}: {entry!r}")


@dataclass
class Ledger:
    source_id: str
    rows: list[dict] = field(default_factory=list)
    filings: dict[str, dict] = field(default_factory=dict)
    reconciliation: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    transactions: list[dict] = field(default_factory=list)

    def filing_id_for(self, source_id: str, Filer_ID, Filer_NamL=None, Report_Num=None, Rpt_Date=None,
                      From_Date=None, Thru_Date=None, **_ignored) -> str:
        return _filing_id(_filing_key(source_id, str(Filer_ID), _text(Filer_NamL), Report_Num, Rpt_Date,
                                      From_Date, Thru_Date))

    def counts(self) -> dict:
        out: dict = {}
        for rec in self.reconciliation:
            group = "with_rows" if rec["rows"] else "summary_only"
            bucket = out.setdefault(rec["schedule"], {"with_rows": {}, "summary_only": {}})[group]
            bucket[rec["status"]] = bucket.get(rec["status"], 0) + 1
        return out


def _readable_columns(sheet: str) -> frozenset[str]:
    """The only columns the ledger ever takes from a sheet; street addresses are never among them."""
    cols = set(SHEET_HEADERS.get(sheet, FILING_HEADERS)) | {"Committee_Type"}
    schedule = SCHEDULES.get(sheet)
    if schedule:
        cols |= {"Tran_Type", f"{schedule.name}_NamT", f"{schedule.name}_NamS", *(c for _, c in schedule.reported)}
    return frozenset(cols)


def _physical_rows(rel: str, path: Path):
    """Yield (sheet, excel_row_number, {header: value}) for every data row, holding only readable columns."""
    with zipfile.ZipFile(path) as zf:
        inner = next(n for n in zf.namelist() if n.endswith(".xlsx"))
        with zf.open(inner) as f:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            try:
                for sheet in wb.sheetnames:
                    rows = wb[sheet].iter_rows(values_only=True)
                    readable = _readable_columns(sheet)
                    keep = [(i, h) for i, h in enumerate(next(rows, ())) if h in readable]
                    for number, values in enumerate(rows, start=2):
                        yield sheet, number, {h: values[i] if i < len(values) else None for i, h in keep}
            finally:
                wb.close()


def _schedule_row(ref: dict, filing_id: str, schedule: Schedule, values: dict) -> dict:
    row = {"row_ref": ref, "filing_id": filing_id, "schedule": schedule.letter, "disposition": "retained",
           "filer_name": _text(values.get("Filer_NamL")), "committee_type": _text(values.get("Committee_Type")),
           "tran_id": _text(values.get("Tran_ID")), "tran_type": _text(values.get("Tran_Type")),
           "entity_cd": _text(values.get("Entity_Cd")), "tran_date": _iso(values.get(schedule.date)),
           "memo_code": _text(values.get("Memo_Code")), "memo_ref": _text(values.get("Memo_RefNo")),
           "name": {part: _text(values.get(f"{schedule.name}_Nam{suffix}"))
                    for part, suffix in (("last", "L"), ("first", "F"), ("title", "T"), ("suffix", "S"))},
           "reported": {key: (_zip5(values.get(col)) if key == "zip5" else _text(values.get(col)))
                        for key, col in schedule.reported}}
    # CAL format: a Memo_Code marks an informational entry outside the schedule's totals. A Memo_RefNo alone
    # only cross-references another entry and the row still counts.
    row["additive"] = row["memo_code"] is None
    if _text(values.get("Form_Type")) != schedule.letter:
        row.update(disposition="rejected", reason="unexpected_form_type")
    elif row["tran_id"] is None:
        row.update(disposition="rejected", reason="no_tran_id")
    else:
        try:
            row["amount"] = str(_cents(values.get(schedule.amount)))
        except ValueError:
            row.update(disposition="rejected", reason="unparseable_amount")
    return row


def build_ledger(source_id: str, workbooks: list[tuple[str, Path]], version_evidence=(), exceptions=()) -> Ledger:
    """Account for every physical row of every workbook and reconcile Schedule A/E per filing."""
    ledger = Ledger(source_id)
    oracles: dict[tuple[str, str], list[Decimal]] = defaultdict(list)
    for rel, path in workbooks:
        for sheet, number, values in _physical_rows(rel, path):
            ref = {"file": rel, "sheet": sheet, "row": number}
            filer_id = _text(values.get("Filer_ID"))
            if filer_id is None:
                schedule = SCHEDULES[sheet].letter if sheet in SCHEDULES else None
                ledger.rows.append({"row_ref": ref, "filing_id": None, "schedule": schedule,
                                    "disposition": "rejected", "reason": "no_filer_id"})
                if schedule:
                    ledger.errors.append(f"{rel}!{sheet}!{number}: no_filer_id, the row belongs to no filing")
                continue
            key = _filing_key(source_id, filer_id, _text(values.get("Filer_NamL")), values.get("Report_Num"),
                              values.get("Rpt_Date"), values.get("From_Date"), values.get("Thru_Date"))
            fid = _filing_id(key)
            filing = ledger.filings.setdefault(fid, {
                "filing_id": fid, "source_id": source_id, "filer_id": filer_id, "filer_key": key[1],
                "from_date": key[2], "thru_date": key[3], "rpt_date": key[4], "report_num": key[5],
                "files": set(), "names": set(), "sheets": set(), "committee_types": set()})
            filing["files"].add(rel)
            filing["names"].add(_text(values.get("Filer_NamL")))
            filing["sheets"].add(sheet)
            if _text(values.get("Committee_Type")):
                filing["committee_types"].add(_text(values.get("Committee_Type")))
            if sheet in SCHEDULES:
                ledger.rows.append(_schedule_row(ref, fid, SCHEDULES[sheet], values))
                continue
            row = {"row_ref": ref, "filing_id": fid, "schedule": None, "disposition": "excluded",
                   "reason": "out_of_scope_sheet"}
            if sheet == "Summary":
                row["reason"] = "summary_line"
                form, line = _text(values.get("Form_Type")), _text(values.get("Line_Item"))
                if form in ("A", "E") and line == ORACLE_LINE:
                    try:
                        oracles[(fid, form)].append(_cents(values.get("Amount_A")))
                    except ValueError:
                        ledger.errors.append(f"{rel}!Summary!{number}: unparseable oracle amount")
            ledger.rows.append(row)
    ledger.rows.sort(key=lambda r: (r["row_ref"]["file"], r["row_ref"]["sheet"], r["row_ref"]["row"]))
    _mark_ambiguity(ledger, oracles)
    _apply_versions(ledger, version_evidence)
    _reconcile(ledger, oracles, exceptions)
    for filing in ledger.filings.values():
        for key in ("files", "names", "sheets", "committee_types"):
            filing[key] = sorted(filing[key], key=lambda v: (v is None, v or ""))
    return ledger


def _mark_ambiguity(ledger: Ledger, oracles: dict) -> None:
    for fid, filing in ledger.filings.items():
        reasons = []
        if len(filing["files"]) > 1:
            reasons.append("multiple_files")
        if len(filing["names"]) > 1:
            reasons.append("multiple_names")
        reasons += [f"multiple_oracles:{s}" for s in ("A", "E") if len(oracles.get((fid, s), ())) > 1]
        filing["ambiguous"] = reasons


def _chain_end(fid: str, superseded_by: dict[str, str]) -> str:
    while fid in superseded_by:
        fid = superseded_by[fid]
    return fid


def _apply_versions(ledger: Ledger, version_evidence) -> None:
    """Supersede only on explicit, located source evidence; otherwise competing versions stay unresolved."""
    families: dict[tuple, list[str]] = defaultdict(list)
    for fid, f in ledger.filings.items():
        f["version_status"] = "single"
        if f["from_date"] and f["thru_date"]:
            families[(f["filer_key"], f["from_date"], f["thru_date"])].append(fid)
    superseded_by: dict[str, str] = {}
    for entry in version_evidence:
        _require(entry, ("original", "amended", "locator", "evidence"), "version evidence")
        original, amended = ledger.filings.get(entry["original"]), ledger.filings.get(entry["amended"])
        if not original or not amended:
            raise LedgerError(f"version evidence names an unknown filing: {entry!r}")
        if entry["original"] == entry["amended"]:
            raise LedgerError(f"version evidence links a filing to itself: {entry!r}")
        if (original["filer_key"], original["from_date"], original["thru_date"]) != \
                (amended["filer_key"], amended["from_date"], amended["thru_date"]):
            raise LedgerError(f"version evidence spans two filing families: {entry!r}")
        if superseded_by.get(entry["original"], entry["amended"]) != entry["amended"]:
            raise LedgerError(f"version evidence gives one filing two successors: {entry!r}")
        superseded_by[entry["original"]] = entry["amended"]
    for fid in superseded_by:  # every chain must end, without a cycle
        seen = {fid}
        while fid in superseded_by:
            fid = superseded_by[fid]
            if fid in seen:
                raise LedgerError(f"version evidence forms a cycle through {fid}")
            seen.add(fid)
    for members in families.values():
        if len(members) < 2:
            continue
        current = [fid for fid in members if fid not in superseded_by]
        resolved = len(current) == 1 and all(_chain_end(fid, superseded_by) == current[0] for fid in members)
        for fid in members:
            if not resolved:
                ledger.filings[fid]["version_status"] = "unresolved_versions"
            elif fid in superseded_by:
                ledger.filings[fid].update(version_status="superseded", superseded_by=superseded_by[fid])
            else:
                ledger.filings[fid]["version_status"] = "current"
    for row in ledger.rows:
        filing = ledger.filings.get(row["filing_id"]) if row["filing_id"] else None
        if filing and filing["version_status"] == "superseded" and row["disposition"] == "retained":
            row.update(disposition="superseded", superseded_by=filing["superseded_by"])


def _reconcile(ledger: Ledger, oracles: dict, exceptions) -> None:
    excepted = {}
    for entry in exceptions:
        _require(entry, ("filing_id", "schedule", "locator", "evidence"), "reconciliation exception")
        excepted[(entry["filing_id"], entry["schedule"])] = entry
    groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in ledger.rows:
        if row["schedule"] and row["filing_id"]:
            groups[(row["filing_id"], row["schedule"])].append(row)
    for fid, schedule in sorted(set(groups) | set(oracles)):
        rows, filing = groups.get((fid, schedule), []), ledger.filings[fid]
        itemized = sum((Decimal(r["amount"]) for r in rows
                        if r["disposition"] in ("retained", "superseded") and r["additive"]), Decimal("0.00"))
        found = oracles.get((fid, schedule), [])
        oracle = found[0] if len(found) == 1 else None
        if filing["version_status"] in ("superseded", "unresolved_versions"):
            status = filing["version_status"]
        elif filing["ambiguous"]:
            status = "ambiguous"
        elif oracle is None:
            status = "missing_oracle"
        elif itemized == oracle:
            status = "matched"
        elif (fid, schedule) in excepted:
            status = "excepted"
        else:
            status = "mismatched"
            ledger.errors.append(f"{fid} schedule {schedule} mismatched: itemized {itemized} vs oracle {oracle}")
        rejected = sum(r["disposition"] == "rejected" for r in rows)
        ledger.reconciliation.append({
            "filing_id": fid, "schedule": schedule, "rows": len(rows), "rejected_rows": rejected,
            "itemized": str(itemized), "oracle": None if oracle is None else str(oracle), "status": status,
            "validated": status == "matched" and rejected == 0})
    for key in set(excepted) - {(r["filing_id"], r["schedule"]) for r in ledger.reconciliation
                                if r["status"] == "excepted"}:
        ledger.errors.append(f"reconciliation exception {key} does not match a mismatched group")


def _dump(obj) -> str:
    return json.dumps(obj, sort_keys=True, ensure_ascii=False)


def write_ledger(ledger: Ledger, out_dir: Path) -> dict:
    """Write ledger.jsonl, filings.jsonl and reconciliation.json: sorted, no timestamps, byte-stable."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "ledger.jsonl", "w") as f:
        for row in ledger.rows:
            f.write(_dump(row) + "\n")
    with open(out_dir / "filings.jsonl", "w") as f:
        for fid in sorted(ledger.filings):
            f.write(_dump(ledger.filings[fid]) + "\n")
    dispositions: dict[str, int] = {}
    for row in ledger.rows:
        dispositions[row["disposition"]] = dispositions.get(row["disposition"], 0) + 1
    report = {
        "source_id": ledger.source_id,
        "physical_rows": len(ledger.rows),
        "filings": len(ledger.filings),
        "dispositions": dispositions,
        "counts": ledger.counts(),
        "groups": sorted(ledger.reconciliation, key=lambda r: (r["filing_id"], r["schedule"])),
        "errors": ledger.errors,
    }
    (out_dir / "reconciliation.json").write_text(json.dumps(report, indent=2, sort_keys=True) + "\n")
    return report


# ---------------------------------------------------------------------------
# Transactions: one per real-world payment, however many filings report it
# ---------------------------------------------------------------------------

def _overlap(a: dict, b: dict) -> bool:
    """Two filings' periods overlap; a filing without a period may overlap anything."""
    if not (a["from_date"] and a["thru_date"] and b["from_date"] and b["thru_date"]):
        return True
    return a["from_date"] <= b["thru_date"] and b["from_date"] <= a["thru_date"]


def _content(row: dict) -> tuple:
    return (row["amount"], row["tran_date"], tuple(row["name"].values()), row["entity_cd"])


def _split_repeats(rows: list[dict], filings: dict) -> list[tuple[str, list[dict]]]:
    """Split one (filer, schedule, Tran_ID) group into transactions, each tagged with its kind."""
    per_filing = defaultdict(list)
    for row in rows:
        per_filing[row["filing_id"]].append(row)
    if len(per_filing) == 1:
        kind = "single" if len(rows) == 1 else "repeated_within_filing"  # the filer's own total counts each
        return [(kind, [row]) for row in rows]
    contents = {_content(r) for r in rows}
    if max(len(v) for v in per_filing.values()) == 1 and len(contents) == 1:
        return [("duplicate_report", rows)]
    ids = sorted(per_filing)
    if all(not _overlap(filings[a], filings[b]) for i, a in enumerate(ids) for b in ids[i + 1:]):
        return [("reused_tran_id", [row]) for row in rows]
    return [("competing_versions", rows)]


def build_transactions(ledger: Ledger) -> list[dict]:
    """Group retained A/E rows into transactions, decide which count, and bridge every row to its flow.

    A transaction counts (and becomes a public MoneyFlow) only when every reporting row is additive and
    nonzero, every reporting filing's schedule is validated, and no competing version of it exists.
    """
    validated = {(g["filing_id"], g["schedule"]): g["validated"] for g in ledger.reconciliation}
    order = {fid: (f["rpt_date"] or "", f["report_num"] or "", fid) for fid, f in ledger.filings.items()}
    groups = defaultdict(list)
    for row in ledger.rows:
        if row["schedule"] and row["disposition"] == "retained":
            filing = ledger.filings[row["filing_id"]]
            groups[(filing["filer_key"], row["schedule"], row["tran_id"])].append(row)
    txs = []
    for (_, schedule, tran_id), rows in sorted(groups.items()):
        rows.sort(key=lambda r: (order[r["filing_id"]], r["row_ref"]["file"], r["row_ref"]["row"]))
        for kind, members in _split_repeats(rows, ledger.filings):
            primary = members[0]
            if not all(validated.get((r["filing_id"], schedule)) for r in members):
                reason = "filing_not_validated"
            elif kind == "competing_versions":
                reason = "competing_versions"
            elif not all(r["additive"] for r in members):
                reason = "non_additive"
            elif Decimal(primary["amount"]) == 0:
                reason = "zero_amount"
            else:
                reason = None
            filer_id = ledger.filings[primary["filing_id"]]["filer_id"]
            txs.append({
                "kind": kind, "schedule": schedule, "filer_id": filer_id, "tran_id": tran_id,
                "amount": primary["amount"], "tran_date": primary["tran_date"], "name": primary["name"],
                "entity_cd": primary["entity_cd"], "filing_ids": sorted({r["filing_id"] for r in members}),
                "rows": [r["row_ref"] for r in members], "counts": reason is None, "reason": reason,
                "base_id": f"moneyflow-{filer_id}-{tran_id}",
                "_members": members})
    _assign_ids(txs)
    for tx in txs:
        for i, row in enumerate(tx.pop("_members")):
            row.update(moneyflow_id=tx["moneyflow_id"], counted=tx["counts"], withheld_reason=tx["reason"],
                       transaction_kind=tx["kind"], primary=i == 0)
    _bridge(ledger)
    ledger.transactions = txs
    return txs


def _assign_ids(txs: list[dict]) -> None:
    """Keep the live `moneyflow-{Filer_ID}-{Tran_ID}` id unless it would name more than one transaction."""
    per_base = defaultdict(int)
    for tx in txs:
        per_base[tx["base_id"]] += 1
    occurrences: dict[str, int] = defaultdict(int)
    for tx in txs:
        if per_base[tx["base_id"]] == 1:
            tx["moneyflow_id"] = tx["base_id"]
            continue
        key = [tx["filing_ids"][0], tx["amount"], tx["tran_date"], list(tx["name"].values()), tx["entity_cd"]]
        digest = hashlib.sha256(json.dumps(key).encode()).hexdigest()[:8]
        tx["moneyflow_id"] = f"{tx['base_id']}-{tx['schedule'].lower()}-{digest}"
        occurrences[tx["moneyflow_id"]] += 1
        if occurrences[tx["moneyflow_id"]] > 1:  # identical twins in one filing: number them in export order
            tx["moneyflow_id"] += f"-{occurrences[tx['moneyflow_id']]}"
    ids = [tx["moneyflow_id"] for tx in txs]
    if len(ids) != len(set(ids)):
        raise LedgerError("two transactions share a MoneyFlow id")


def _bridge(ledger: Ledger) -> None:
    """Explain each filing's itemized total as emitted + counted elsewhere + superseded + withheld."""
    buckets = defaultdict(lambda: {"emitted": Decimal(0), "counted_in_other_filing": Decimal(0),
                                   "superseded": Decimal(0), "non_additive": Decimal(0), "withheld": {}})
    for row in ledger.rows:
        if not row["schedule"] or not row["filing_id"] or row["disposition"] == "rejected":
            continue
        b = buckets[(row["filing_id"], row["schedule"])]
        amount = Decimal(row["amount"])
        if not row["additive"]:
            b["non_additive"] += amount
        elif row["disposition"] == "superseded":
            b["superseded"] += amount
        elif row["counted"]:
            b["emitted" if row["primary"] else "counted_in_other_filing"] += amount
        else:
            reason = row["withheld_reason"]
            b["withheld"][reason] = b["withheld"].get(reason, Decimal(0)) + amount
    for group in ledger.reconciliation:
        b = buckets[(group["filing_id"], group["schedule"])]
        explained = b["emitted"] + b["counted_in_other_filing"] + b["superseded"] + sum(b["withheld"].values())
        if explained != Decimal(group["itemized"]):
            raise LedgerError(f"bridge does not explain {group['filing_id']} {group['schedule']}")
        group["bridge"] = {k: (str(v.quantize(CENT)) if isinstance(v, Decimal)
                               else {r: str(a.quantize(CENT)) for r, a in sorted(v.items())})
                           for k, v in b.items()}
